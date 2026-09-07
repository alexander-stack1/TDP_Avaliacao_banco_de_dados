#!/usr/bin/env python3
"""Gera docs/Entrega_Case_Bolsa_Valores.docx com os modelos conceitual e lógico
(o arquivo .SQL vai anexado à parte, conforme a orientação da atividade).

Uso: python3 scripts/gerar_docx.py      (requer: pip install python-docx)
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

RAIZ = Path(__file__).resolve().parent.parent
SAIDA = RAIZ / "docs/Entrega_Case_Bolsa_Valores.docx"
PNG_CONC = RAIZ / "docs/01-modelo-conceitual/modelo_conceitual.png"
PNG_LOG = RAIZ / "docs/02-modelo-logico/modelo_logico.png"

doc = Document()
estilo = doc.styles["Normal"]
estilo.font.name = "Calibri"
estilo.font.size = Pt(11)

sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width
for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, lado, Cm(1.8))
LARGURA_UTIL = sec.page_width - sec.left_margin - sec.right_margin


def tabela(cabecalho, linhas, larguras=None):
    t = doc.add_table(rows=1, cols=len(cabecalho))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(cabecalho):
        c = t.rows[0].cells[i]
        c.text = h
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9.5)
    for linha in linhas:
        cells = t.add_row().cells
        for i, v in enumerate(linha):
            cells[i].text = str(v)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
    if larguras:
        for row in t.rows:
            for i, w in enumerate(larguras):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


# --------------------------------------------------------------------- capa
titulo = doc.add_heading("Case — Negociações na Bolsa de Valores", level=0)
titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph("Modelagem de Banco de Dados: modelos conceitual, lógico e físico")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph("SGBD alvo: PostgreSQL 16 · Ferramenta do conceitual: BR Modelo (notação Chen)")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()
doc.add_paragraph("Equipe: ______________________________________________________________")
doc.add_paragraph("Data: ____ / ____ / ________")
doc.add_paragraph()
doc.add_paragraph(
    "Anexos desta entrega: (1) este documento, com os modelos conceitual e lógico; "
    "(2) o arquivo 00_bolsa_completo.sql com DDL, DML e DQL; "
    "(3) o arquivo modelo_conceitual.brmodelo.json (brModelo Web)."
)

# --------------------------------------------------------------------- 1. enunciado
doc.add_page_break()
doc.add_heading("1. Descrição do problema", level=1)
doc.add_paragraph(
    "Uma corretora deseja gerenciar Investidores, Ações e suas Negociações na bolsa de valores. "
    "Cada investidor é identificado por CPF ou CNPJ e possui nome completo, tipo (pessoa física ou jurídica), "
    "e-mail e telefone. Cada ação pertence a uma empresa listada, identificada pelo ticker, com nome da empresa, "
    "setor de atuação e valor de mercado. Os investidores realizam negociações (compra ou venda) com data e hora, "
    "tipo de operação, quantidade e valor unitário. A corretora mantém o histórico de cotações (data, hora e valor "
    "por ação) para análises de séries temporais e acompanha o saldo de carteira de cada investidor, atualizado a "
    "partir das negociações e passível de análise retrospectiva com o histórico de cotações."
)

doc.add_heading("1.1 Requisitos identificados", level=2)
for r in [
    "RF01 — Cadastrar investidores PF (CPF) e PJ (CNPJ), sem duplicidade de documento.",
    "RF02 — Cadastrar empresas listadas e seus papéis (uma empresa pode ter mais de um ticker).",
    "RF03 — Registrar negociações de compra e venda com data/hora, quantidade e preço do momento.",
    "RF04 — Armazenar o histórico intradiário de cotações por ação.",
    "RF05 — Manter a posição (carteira) de cada investidor atualizada a partir das negociações.",
    "RF06 — Permitir análise retrospectiva: valor da carteira em qualquer instante passado.",
    "RN01 — Um investidor não pode vender mais ações do que possui.",
    "RN02 — Negociação é registro contábil: não pode ser alterada nem excluída.",
    "RN03 — Só existe uma cotação por ação em cada instante.",
]:
    doc.add_paragraph(r, style="List Bullet")

# --------------------------------------------------------------------- 2. conceitual
doc.add_page_break()
doc.add_heading("2. Modelo Conceitual (BR Modelo — notação Chen)", level=1)
doc.add_picture(str(PNG_CONC), width=LARGURA_UTIL)
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph("Figura 1 — Diagrama Entidade-Relacionamento (arquivo: modelo_conceitual.brmodelo.json).").alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("2.1 Entidades e atributos", level=2)
tabela(["Entidade", "Descrição", "Atributos (identificador em MAIÚSCULAS)"], [
    ("EMPRESA", "Companhia listada na bolsa", "ID_EMPRESA, cnpj, nome, setor, valor_mercado"),
    ("AÇÃO", "Papel negociado, identificado pelo ticker", "ID_ACAO, ticker, tipo_acao, ativa"),
    ("INVESTIDOR", "Cliente da corretora (PF ou PJ)",
     "ID_INVESTIDOR, documento (CPF/CNPJ), tipo_investidor, nome_completo, email, telefone"),
    ("NEGOCIAÇÃO", "Uma compra ou venda de uma ação por um investidor",
     "ID_NEGOCIACAO, data_hora, tipo_operacao, quantidade, valor_unitario"),
    ("COTAÇÃO (fraca)", "Preço de uma ação em um instante", "data_hora (chave parcial), valor"),
], larguras=[4, 8, 12])

doc.add_heading("2.2 Relacionamentos e cardinalidades", level=2)
tabela(["Relacionamento", "Entidades", "Cardinalidade (mín,máx)", "Atributos"], [
    ("emite", "EMPRESA — AÇÃO", "EMPRESA (1,n) · AÇÃO (1,1)", "—"),
    ("possui (identificador)", "AÇÃO — COTAÇÃO", "AÇÃO (1,n) · COTAÇÃO (1,1)", "—"),
    ("realiza", "INVESTIDOR — NEGOCIAÇÃO", "INVESTIDOR (0,n) · NEGOCIAÇÃO (1,1)", "—"),
    ("refere-se a", "AÇÃO — NEGOCIAÇÃO", "AÇÃO (0,n) · NEGOCIAÇÃO (1,1)", "—"),
    ("mantém (carteira)", "INVESTIDOR — AÇÃO", "INVESTIDOR (0,n) · AÇÃO (0,n)", "quantidade, preco_medio"),
], larguras=[5, 6, 8, 5])

doc.add_heading("2.3 Decisões de modelagem", level=2)
for d in [
    "Negociação é entidade, e não um relacionamento N:N entre Investidor e Ação: o mesmo par ocorre muitas vezes, "
    "cada uma com data/hora, tipo, quantidade e preço próprios. Um relacionamento N:N admitiria uma única ocorrência por par.",
    "Carteira é o relacionamento N:N “mantém”, com atributos: a posição atual é única por par (investidor, ação) "
    "e é derivada das negociações — regra implementada por trigger no modelo físico.",
    "Cotação é entidade fraca de Ação, identificada por (ação, data_hora); o relacionamento “possui” é identificador.",
    "Empresa foi separada de Ação: nome, setor e valor de mercado pertencem à companhia, que pode emitir mais de um papel "
    "(ex.: PETR3 e PETR4).",
    "Investidor guarda CPF ou CNPJ num único atributo “documento”, qualificado por “tipo_investidor”; no físico isso vira "
    "UNIQUE + CHECK de 11 ou 14 dígitos.",
]:
    doc.add_paragraph(d, style="List Number")

# --------------------------------------------------------------------- 3. lógico
doc.add_page_break()
doc.add_heading("3. Modelo Lógico Relacional", level=1)
doc.add_picture(str(PNG_LOG), width=LARGURA_UTIL)
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph("Figura 2 — Esquema relacional (esquema bolsa).").alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("3.1 Notação textual", level=2)
for linha in [
    "EMPRESA (id_empresa, cnpj, nome, setor, valor_mercado, atualizado_em)  —  PK id_empresa · UQ cnpj",
    "ACAO (id_acao, ticker, id_empresa, tipo_acao, ativa)  —  PK id_acao · UQ ticker · FK id_empresa → EMPRESA",
    "INVESTIDOR (id_investidor, documento, tipo_investidor, nome_completo, email, telefone, criado_em)  —  PK id_investidor · UQ documento · UQ email",
    "COTACAO (id_cotacao, id_acao, data_hora, valor)  —  PK id_cotacao · UQ (id_acao, data_hora) · FK id_acao → ACAO",
    "NEGOCIACAO (id_negociacao, id_investidor, id_acao, data_hora, tipo_operacao, quantidade, valor_unitario, valor_total)  —  PK id_negociacao · FK id_investidor → INVESTIDOR · FK id_acao → ACAO",
    "CARTEIRA (id_investidor, id_acao, quantidade, preco_medio, atualizado_em)  —  PK (id_investidor, id_acao) · FKs → INVESTIDOR, ACAO",
]:
    p = doc.add_paragraph()
    r = p.add_run(linha)
    r.font.name = "Consolas"
    r.font.size = Pt(9)

doc.add_heading("3.2 Mapeamento conceitual → lógico", level=2)
tabela(["Elemento conceitual", "Resultado lógico", "Regra"], [
    ("Entidades EMPRESA, AÇÃO, INVESTIDOR, NEGOCIAÇÃO", "Uma tabela cada, PK substituta id_*", "Entidade → tabela"),
    ("emite (1:N)", "FK acao.id_empresa", "1:N → FK no lado N"),
    ("COTAÇÃO (fraca) + possui", "Tabela cotacao com FK id_acao e UNIQUE (id_acao, data_hora)", "Entidade fraca herda a chave do dono"),
    ("realiza / refere-se a (1:N)", "FKs negociacao.id_investidor e negociacao.id_acao", "1:N → FK no lado N"),
    ("mantém (N:N com atributos)", "Tabela carteira, PK composta (id_investidor, id_acao) + atributos", "N:N → tabela associativa"),
    ("Atributo derivado valor_total", "Coluna gerada quantidade × valor_unitario", "Derivado → calculado no SGBD"),
], larguras=[7, 10, 7])

doc.add_heading("3.3 Dicionário de dados", level=2)
dicionario = {
    "EMPRESA": [
        ("id_empresa", "bigint identity", "PK"), ("cnpj", "text", "NOT NULL, UNIQUE, 14 dígitos"),
        ("nome", "text", "NOT NULL"), ("setor", "text", "NOT NULL"),
        ("valor_mercado", "numeric(18,2)", "NOT NULL, ≥ 0"), ("atualizado_em", "timestamptz", "NOT NULL, default now()")],
    "ACAO": [
        ("id_acao", "bigint identity", "PK"), ("ticker", "text", "NOT NULL, UNIQUE, 4 letras + 1–2 dígitos"),
        ("id_empresa", "bigint", "NOT NULL, FK → empresa (RESTRICT)"), ("tipo_acao", "text", "IN (ON, PN, UNIT)"),
        ("ativa", "boolean", "NOT NULL, default true")],
    "INVESTIDOR": [
        ("id_investidor", "bigint identity", "PK"),
        ("documento", "text", "NOT NULL, UNIQUE, 11 dígitos (PF) ou 14 (PJ)"),
        ("tipo_investidor", "text", "IN (PF, PJ)"), ("nome_completo", "text", "NOT NULL"),
        ("email", "text", "NOT NULL, UNIQUE, formato válido"), ("telefone", "text", "NULL, 10–11 dígitos"),
        ("criado_em", "timestamptz", "NOT NULL, default now()")],
    "COTACAO": [
        ("id_cotacao", "bigint identity", "PK"), ("id_acao", "bigint", "NOT NULL, FK → acao (CASCADE)"),
        ("data_hora", "timestamptz", "NOT NULL, UNIQUE com id_acao"), ("valor", "numeric(12,4)", "NOT NULL, > 0")],
    "NEGOCIACAO": [
        ("id_negociacao", "bigint identity", "PK"), ("id_investidor", "bigint", "NOT NULL, FK → investidor (RESTRICT)"),
        ("id_acao", "bigint", "NOT NULL, FK → acao (RESTRICT)"), ("data_hora", "timestamptz", "NOT NULL, default now()"),
        ("tipo_operacao", "text", "IN (COMPRA, VENDA)"), ("quantidade", "integer", "NOT NULL, > 0"),
        ("valor_unitario", "numeric(12,4)", "NOT NULL, > 0"), ("valor_total", "numeric(18,2)", "gerada: quantidade × valor_unitario")],
    "CARTEIRA": [
        ("id_investidor", "bigint", "PK, FK → investidor (CASCADE)"), ("id_acao", "bigint", "PK, FK → acao (RESTRICT)"),
        ("quantidade", "integer", "NOT NULL, ≥ 0"), ("preco_medio", "numeric(12,4)", "NOT NULL, ≥ 0"),
        ("atualizado_em", "timestamptz", "NOT NULL")],
}
for nome, cols in dicionario.items():
    doc.add_paragraph(nome, style="Heading 3")
    tabela(["Coluna", "Tipo", "Restrições"], cols, larguras=[5, 5, 14])

doc.add_heading("3.4 Normalização", level=2)
for n in [
    "1FN — atributos atômicos (um telefone, um documento; nenhuma lista em coluna).",
    "2FN — a única PK composta é a de CARTEIRA; quantidade, preco_medio e atualizado_em dependem do par inteiro.",
    "3FN — nenhum atributo depende de outro não-chave: dados da companhia ficam em EMPRESA, não repetidos por papel; "
    "valor_total é derivado, mas materializado como coluna gerada pelo SGBD (nunca fica inconsistente).",
]:
    doc.add_paragraph(n, style="List Bullet")

# --------------------------------------------------------------------- 4. físico (resumo)
doc.add_page_break()
doc.add_heading("4. Modelo Físico (resumo — o script completo está no arquivo .SQL anexo)", level=1)
doc.add_paragraph(
    "Arquivo: 00_bolsa_completo.sql (PostgreSQL 16). Execução: "
    "psql -v ON_ERROR_STOP=1 -d bolsa_valores -f 00_bolsa_completo.sql"
)
tabela(["Bloco", "Conteúdo"], [
    ("DDL", "Esquema bolsa; 6 tabelas com PK identity, UNIQUE nas chaves naturais, CHECKs (CPF/CNPJ por tipo, ticker, "
            "e-mail, valores positivos), FKs com ação explícita, índices em todas as FKs, índices compostos para extratos "
            "e BRIN na série temporal; coluna gerada valor_total; trigger que atualiza a carteira e rejeita venda sem saldo; "
            "trigger que torna a negociação imutável; função fn_carteira_em para análise retrospectiva; 3 views; "
            "papel bolsa_leitura somente-leitura; COMMENT ON nos objetos."),
    ("DML", "6 investidores (4 PF, 2 PJ), 7 empresas, 8 ações, 600 cotações intradiárias geradas com generate_series, "
            "15 negociações em ordem cronológica (o trigger monta a carteira), UPDATE, DELETE e 4 blocos DO que "
            "demonstram as regras de integridade."),
    ("DQL", "16 consultas: carteira valorizada, patrimônio por investidor, extrato, volume por ação, OHLC diário, "
            "variação e média móvel (janelas), máximas/mínimas, carteira em data passada, evolução do patrimônio, "
            "resultado realizado, ranking, exposição por setor, ações sem negociação, PF × PJ, conferência de integridade "
            "e plano de execução."),
], larguras=[3, 21])
doc.add_paragraph(
    "Validação: o script foi executado integralmente em PostgreSQL 16 sem erros; as regras de integridade "
    "demonstradas retornaram OK e a conferência carteira × negociações retornou zero divergências."
)

SAIDA.parent.mkdir(parents=True, exist_ok=True)
doc.save(SAIDA)
print("gerado:", SAIDA)
